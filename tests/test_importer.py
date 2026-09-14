"""Adversarial fixtures are created under this project's .cache only."""
from argparse import Namespace
from contextlib import redirect_stdout
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel
from datetime import datetime

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / 'scripts'))
import import_excel as importer
from stock_import.normalize import iso_date, missing_dates, number
from stock_import.reader import summary_date, read_stocks


class ImporterTests(unittest.TestCase):
    def setUp(self):
        (PROJECT / '.cache').mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=PROJECT / '.cache', prefix='import-test-')
        self.root = Path(self.temporary.name)
        self.source = self.root / 'source'
        self.source.mkdir()
        self.data = self.root / 'public' / 'data'
        self.patches = [patch.object(importer, 'ROOT', self.root), patch.object(importer, 'DATA', self.data), patch.object(importer, 'CACHE', self.root / '.cache/state.json'), patch.dict(os.environ, {'STOCK_DATA_SOURCE': str(self.source)})]
        for item in self.patches:
            item.start()

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temporary.cleanup()

    def workbook(self, filename, date='2026-09-10', volume=123456, mtime=1700000000, weekly=False):
        book = Workbook()
        s = book.active
        s.title = '汇总'
        s.append(['标题'])
        s.append(['运行项目', '测试任务', None, '信号', '数量'])
        s.append(['T0日期', datetime.fromisoformat(date), None, '严格两连阴', 1])
        s.append(['异常合计', 0, None, '严格三连阴', 0])
        s.append([None, None, None, '三连阳及以上', 0])
        for name in ['严格两连阴', '严格三连阴', '三连阳及以上']:
            s = book.create_sheet(name)
            s.append(['标题（错误的9999只）'])
            s.append([])
            headers = ['证券简称', '未知新增列', '证券代码', '连续天数', '成交量', '上市日期', 'ST/*ST', '申万一级行业']
            if weekly:
                headers[1:1] = ['周RSI14', '周KDJ-J', '周KDJ-D', '周KDJ-K']
            s.append(headers)
            if name == '严格两连阴':
                values = ['样本', '忽略', '000607.SZ', 2, volume, to_excel(datetime(1996,8,30)), None, '传媒']
                if weekly:
                    values[1:1] = [41.2876, -5.25, 38.302, 38.978]
                s.append(values)
        s = book.create_sheet('异常'); s.append(['异常类型', '证券代码'])
        s = book.create_sheet('行业汇总'); s.append(['类别', '排名', '申万一级行业', '数量']);s.append(['严格两连阴',1,'传媒',1])
        path=self.source/filename;book.save(path);book.close();os.utime(path,(mtime,mtime));return path

    def run_import(self, date=None, rebuild=False):
        with redirect_stdout(io.StringIO()) as output:
            code=importer.run(Namespace(date=date,rebuild=rebuild))
        return code,output.getvalue()

    def index(self):
        return json.loads((self.data/'index.json').read_text(encoding='utf-8'))

    def test_dates_and_normalization(self):
        self.assertEqual(iso_date('1996-08-30'),'1996-08-30')
        self.assertEqual(iso_date('20260910'),'2026-09-10')
        self.assertEqual(iso_date(to_excel(datetime(2026,9,10))),'2026-09-10')
        self.assertEqual(missing_dates('20260904、20260903',datetime(1899,12,30)),['2026-09-04','2026-09-03'])
        self.assertIsNone(number(float('nan')))

    def test_real_header_reorder_null_dates_counts_and_source_unchanged(self):
        path=self.workbook('wrong-filename.xlsx');before=hashlib.sha256(path.read_bytes()).hexdigest()
        code,_=self.run_import();self.assertEqual(code,0)
        day=self.index()['days'][0];self.assertEqual(day['date'],'2026-09-10');self.assertEqual(day['counts']['two-yin'],1)
        record=json.loads((self.data/day['files']['two-yin']).read_text(encoding='utf-8'))[0]
        self.assertEqual(record['code'],'000607.SZ');self.assertEqual(record['listedDate'],'1996-08-30');self.assertIsNone(record['isST']);self.assertIsNone(record['open']);self.assertIsInstance(record['volume'],int)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
        before=(self.data/'index.json').read_bytes();_,log=self.run_import();self.assertIn('没有发现新的交易日数据',log);self.assertEqual((self.data/'index.json').read_bytes(),before)

    def test_duplicate_latest_file_wins_and_history_preserved(self):
        self.workbook('first.xlsx');self.workbook('newest.xlsx',volume=999,mtime=1700000020);self.workbook('history.xlsx',date='2026-09-09')
        _,log=self.run_import();self.assertIn('发现同日期多个统计文件',log)
        day=self.index()['days'][0];self.assertEqual(day['sourceFile'],'newest.xlsx');oldpath=self.data/day['summaryPath']
        self.run_import(date='2026-09-10',rebuild=True)
        self.assertEqual(len(self.index()['dates']),2);self.assertTrue(oldpath.exists())
        summary=json.loads((self.data/self.index()['days'][0]['summaryPath']).read_text(encoding='utf-8'))
        self.assertEqual(len(summary['source']['candidates']),2)

    def test_corrupt_file_and_temporary_file_isolation(self):
        self.workbook('valid.xlsx');(self.source/'broken.xlsx').write_text('not an xlsx');(self.source/'~$locked.xlsx').write_text('ignored')
        os.utime(self.source/'broken.xlsx',(1700000000,1700000000))
        code,log=self.run_import();self.assertEqual(code,2);self.assertIn('扫描文件数量=2',log);self.assertEqual(len(self.index()['days']),1)

    def test_missing_and_empty_sheet_do_not_fake_zero_counts(self):
        path=self.workbook('missing.xlsx')
        from openpyxl import load_workbook
        book=load_workbook(path);del book['严格三连阴'];book['三连阳及以上'].delete_rows(1,book['三连阳及以上'].max_row);book.save(path);book.close();os.utime(path,(1700000000,1700000000))
        self.run_import();counts=self.index()['days'][0]['counts'];self.assertIsNone(counts['three-yin']);self.assertIsNone(counts['three-yang-plus']);self.assertEqual(counts['two-yin'],1)

    def test_invalid_t0_never_falls_back_to_filename(self):
        with self.assertRaises(ValueError):summary_date([['T0日期','bad'],['A股连阴连阳扫描（2026-09-10）']],datetime(1899,12,30),[])

    def test_weekly_headers_are_optional_numeric_and_capability_depends_on_valid_values(self):
        self.workbook('old.xlsx', date='2026-09-11')
        self.workbook('new.xlsx', date='2026-09-14', weekly=True)
        code,_=self.run_import();self.assertEqual(code,0)
        days={day['date']:day for day in self.index()['days']}
        old=json.loads((self.data/days['2026-09-11']['files']['two-yin']).read_text(encoding='utf-8'))[0]
        new=json.loads((self.data/days['2026-09-14']['files']['two-yin']).read_text(encoding='utf-8'))[0]
        self.assertIsNone(old['weeklyKdjJ']);self.assertIsNone(old['weeklyRsi14'])
        self.assertFalse(days['2026-09-11']['capabilities']['weeklyTechnicalIndicators'])
        self.assertEqual(new['weeklyKdjK'],38.978)
        self.assertEqual(new['weeklyKdjD'],38.302)
        self.assertEqual(new['weeklyKdjJ'],-5.25)
        self.assertEqual(new['weeklyRsi14'],41.2876)
        self.assertIsInstance(new['weeklyRsi14'],float)
        self.assertTrue(days['2026-09-14']['capabilities']['weeklyTechnicalIndicators'])

    def test_weekly_nulls_do_not_become_zero_and_wholly_missing_day_is_unavailable(self):
        path=self.workbook('partial.xlsx',date='2026-09-14',weekly=True)
        from openpyxl import load_workbook
        book=load_workbook(path)
        sheet=book['严格两连阴']
        sheet.append(['缺失样本',None,None,None,None,None,'000608.SZ',2,100,None,'否','传媒'])
        book.save(path);book.close();os.utime(path,(1700000000,1700000000))
        self.run_import()
        day=self.index()['days'][0]
        rows=json.loads((self.data/day['files']['two-yin']).read_text(encoding='utf-8'))
        self.assertTrue(day['capabilities']['weeklyTechnicalIndicators'])
        self.assertIsNone(rows[1]['weeklyKdjJ']);self.assertIsNone(rows[1]['weeklyRsi14'])
        self.assertNotEqual(rows[1]['weeklyKdjJ'],0)
        book=load_workbook(path)
        sheet=book['严格两连阴']
        for col in range(2,6):sheet.cell(4,col).value=None
        book.save(path);book.close();os.utime(path,(1700000002,1700000002))
        self.run_import(rebuild=True)
        self.assertFalse(self.index()['days'][0]['capabilities']['weeklyTechnicalIndicators'])


if __name__=='__main__':unittest.main()
