"""Incremental read-only Excel -> immutable daily JSON + compact date index."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import os
from pathlib import Path
import sys
import time

from dotenv import load_dotenv
from openpyxl import load_workbook
from stock_import.reader import SIGNALS, rows_of, summary_date, read_summary, read_stocks, read_industry
from stock_import.validate import validate
from stock_import.writer import read_json, write_json

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'public' / 'data'
CACHE = ROOT / '.cache' / 'import-state.json'
VERSION = 1


def now():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')


def log(level, message):
    print(f'[{level}] {message}', flush=True)


def fingerprint(path):
    stat = path.stat()
    return {'mtimeNs': stat.st_mtime_ns, 'size': stat.st_size}


def build_day(book, file, candidates, origin, date, issues, sha):
    summary = read_summary(rows_of(book, '汇总'), book.epoch)
    signals, status = {}, {}
    for signal, sheet in SIGNALS.items():
        signals[signal], status[sheet] = read_stocks(book, sheet, issues)
    exceptions, status['异常'] = read_stocks(book, '异常', issues, exception=True)
    industry, status['行业汇总'] = read_industry(book, issues)
    counts = {signal: len(records) if status[SIGNALS[signal]] == 'available' else None for signal, records in signals.items()}
    counts['exceptions'] = len(exceptions) if status['异常'] == 'available' else None
    yang = sorted(signals['three-yang-plus'], key=lambda r: (-(r['streak'] or 0), r['code']))
    summary.update(schemaVersion=VERSION, date=date, importedAt=now(), counts=counts, sheetStatus=status,
                   source={'file': file.name, 'sha256': sha, 'modifiedAt': datetime.fromtimestamp(file.stat().st_mtime).astimezone().isoformat(),
                           'candidates': [p.name for p in candidates], 'dateOrigin': origin},
                   sources={'market': 'iFinD 当日行情及停牌查询', 'calendar': '上海证券交易所公开交易日历'},
                   highlights={'longestYangDays': yang[0]['streak'] if yang else None, 'longestYangStocks': yang[:8],
                               'stSignalCount': sum(r['isST'] is True for rs in signals.values() for r in rs),
                               'unknownSTCount': sum(r['isST'] is None for rs in signals.values() for r in rs)},
                   exceptionTypes=dict(Counter(r['type'] for r in exceptions)))
    summary['validations'] = validate(summary, signals, exceptions, industry, issues)
    return summary, {**signals, 'exceptions': exceptions, 'industry': industry}


def stock_bucket(code):
    value = 2166136261
    for char in code:
        value = ((value ^ ord(char)) * 16777619) & 0xffffffff
    return f'{value % 256:02x}'


def update_stock_history(days, old_index, changed):
    # Only explicit per-stock history requests download these compact shards.
    shards = defaultdict(dict)
    for day in days:
        for signal in (*SIGNALS, 'exceptions'):
            for row in read_json(DATA / day['files'][signal], []):
                code = row['code']
                entry = shards[stock_bucket(code)].setdefault(code, {'code': code, 'name': row['name'], 'records': []})
                entry['records'].append({'date': day['date'], 'signal': signal, 'streak': row['streak'], 'type': row.get('type')})
    paths = {}
    for shard, stocks in shards.items():
        # This is a rebuildable derived index. Fixed 256 buckets bound file growth;
        # daily source versions remain immutable and retain every historical date.
        path = f'stock-history/current/{shard}.json'
        write_json(DATA / path, stocks)
        paths[shard] = path
    return paths


def run(args):
    load_dotenv(ROOT / '.env', encoding='utf-8-sig')
    source_value = os.environ.get('STOCK_DATA_SOURCE')
    if not source_value:
        log('ERROR', '请复制 .env.example 为 .env 并设置 STOCK_DATA_SOURCE')
        return 1
    source = Path(source_value).expanduser().resolve()
    if not source.is_dir():
        log('ERROR', f'数据目录不存在：{source}')
        return 1
    if DATA.resolve().is_relative_to(source) or source.is_relative_to(DATA.resolve()):
        log('ERROR', '原始数据目录与生成数据目录不能互相包含')
        return 1
    files = sorted((p for p in source.iterdir() if p.suffix.lower() == '.xlsx' and not p.name.startswith('~$')), key=lambda p: p.name)
    state = read_json(CACHE, {'version': VERSION, 'files': {}})
    old_index = read_json(DATA / 'index.json', {'days': [], 'dates': [], 'stockHistoryShards': {}})
    existing = {day['date']: day for day in old_index['days']}
    groups = defaultdict(list)
    failed = skipped = imported = warnings = 0
    scan_state = {}
    for file in files:
        try:
            mark = fingerprint(file)
            if time.time() - file.stat().st_mtime < 2:
                raise ValueError('文件刚刚被修改，可能正在写入；请稍后重新运行')
            cached = state['files'].get(file.name, {})
            if not args.rebuild and cached.get('fingerprint') == mark and cached.get('date'):
                date, origin, issues = cached['date'], cached['origin'], cached.get('dateIssues', [])
            else:
                issues = []
                book = load_workbook(file, read_only=True, data_only=True)
                try:
                    date, origin = summary_date(rows_of(book, '汇总'), book.epoch, issues)
                finally:
                    book.close()
                if mark != fingerprint(file):
                    raise ValueError('扫描期间文件发生变化，请稍后重试')
            scan_state[file.name] = {'fingerprint': mark, 'date': date, 'origin': origin, 'dateIssues': issues}
            groups[date].append(file)
        except Exception as error:
            failed += 1
            log('ERROR', f'{file.name}: {error}；保留已发布历史数据')
    changed = []
    for date, candidates in sorted(groups.items()):
        candidates.sort(key=lambda p: (p.stat().st_mtime_ns, p.name), reverse=True)
        selected = candidates[0]
        if len(candidates) > 1:
            log('WARNING', f'发现同日期多个统计文件 {date}，选择修改时间最新的 {selected.name}')
            warnings += 1
        if args.date and args.date != date:
            skipped += len(candidates)
            continue
        cached = scan_state[selected.name]
        previous = existing.get(date)
        signature = {'file': selected.name, **cached['fingerprint'], 'candidates': [p.name for p in candidates]}
        if not args.rebuild and previous and previous.get('signature') == signature and all((DATA / p).exists() for p in previous['files'].values()) and (DATA / previous['summaryPath']).exists():
            skipped += len(candidates)
            continue
        try:
            before = fingerprint(selected)
            if before != cached['fingerprint']:
                raise ValueError('文件在扫描后已变化，请稍后重试')
            sha = hashlib.sha256(selected.read_bytes()).hexdigest()
            book = load_workbook(selected, read_only=True, data_only=True)
            try:
                actual_date, _ = summary_date(rows_of(book, '汇总'), book.epoch, [])
                if actual_date != date:
                    raise ValueError('正式读取的 T0日期 与扫描结果不一致，请重试')
                issues = list(cached.get('dateIssues', []))
                if len(candidates) > 1:
                    issues.append({'name': '重复日期', 'status': 'warning', 'message': '按文件修改时间选择最新版本', 'expected': None, 'actual': [p.name for p in candidates]})
                summary, payloads = build_day(book, selected, candidates, cached['origin'], date, issues, sha)
            finally:
                book.close()
            if fingerprint(selected) != before:
                raise ValueError('读取期间文件发生变化，本次不发布')
            revision = sha[:12] + '-' + datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')
            base = f'{date}/{revision}'
            paths = {kind: f'{base}/{kind}.json' for kind in payloads}
            for kind, data in payloads.items():
                write_json(DATA / paths[kind], data)
            summary_path = f'{base}/summary.json'
            write_json(DATA / summary_path, summary)
            issues_count = sum(v['status'] != 'pass' for v in summary['validations'])
            warnings += issues_count
            for item in summary['validations']:
                if item['status'] != 'pass':
                    log('WARNING', f'{date} {item["name"]}: {item["message"]} (预期={item["expected"]}, 实际={item["actual"]})')
            existing[date] = {'date': date, 'revision': revision, 'importedAt': summary['importedAt'], 'sourceFile': selected.name,
                              'signature': signature, 'counts': summary['counts'], 'metrics': summary['metrics'],
                              'tradingStatus': summary['tradingStatus'], 'warningCount': issues_count,
                              'summaryPath': summary_path, 'files': paths}
            changed.append(date)
            imported += 1
            skipped += len(candidates)-1
            log('OK', f'{date} 两连阴={summary["counts"]["two-yin"]} 三连阴={summary["counts"]["three-yin"]} 三连阳及以上={summary["counts"]["three-yang-plus"]} 异常={summary["counts"]["exceptions"]}')
        except Exception as error:
            failed += 1
            log('ERROR', f'{selected.name}: {error}；保留该日期上一版本')
    if args.date and args.date not in groups:
        log('WARNING', f'{args.date} 没有可导入的源文件')
        warnings += 1
    if changed or not (DATA / 'index.json').exists() or old_index.get('stockHistoryFormat') != 'fnv1a-256':
        days = sorted(existing.values(), key=lambda d: d['date'], reverse=True)
        shards = update_stock_history(days, old_index, changed)
        write_json(DATA / 'index.json', {'schemaVersion': VERSION, 'latest': days[0]['date'] if days else None,
                                         'generatedAt': now(), 'dates': [d['date'] for d in days], 'days': days,
                                         'stockHistoryFormat': 'fnv1a-256', 'stockHistoryShards': shards})
    else:
        log('OK', '没有发现新的交易日数据（已有日期的修改文件也已检查）。')
    write_json(CACHE, {'version': VERSION, 'files': scan_state})
    report = {'finishedAt': now(), 'scanned': len(files), 'imported': imported, 'skipped': skipped, 'errors': failed, 'warnings': warnings, 'changedDates': changed}
    write_json(ROOT / '.cache' / 'last-import-report.json', report)
    log('OK', f'扫描文件数量={len(files)} 成功导入数量={imported} 跳过数量={skipped} 异常数量={failed} 警告数量={warnings}')
    return 2 if failed else 0


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--rebuild', action='store_true', help='重新处理所有文件，保留旧版本')
    parser.add_argument('--date', help='只更新指定 T0日期 (YYYY-MM-DD)')
    args = parser.parse_args()
    if args.date:
        try:
            datetime.strptime(args.date, '%Y-%m-%d')
        except ValueError:
            parser.error('--date 必须为 YYYY-MM-DD')
    # Prevent simultaneous writers. A crashed process releases this OS-owned lock.
    ROOT.joinpath('.cache').mkdir(exist_ok=True)
    with (ROOT / '.cache' / 'import.lock').open('a+b') as lock:
        lock.seek(0)
        if lock.read(1) == b'':
            lock.write(b'0')
            lock.flush()
        lock.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(lock.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            log('ERROR', '另一个数据导入进程正在运行，请等待完成')
            return 1
        try:
            return run(args)
        except Exception as error:
            log('ERROR', str(error))
            return 1


if __name__ == '__main__':
    sys.exit(main())
